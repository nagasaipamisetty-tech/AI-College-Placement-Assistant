import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Initialize Presentation with 16:9 Widescreen
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Logo Path
LOGO_PATH = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "data", "agentx_logo.png"))

# Color Palette Constants
COLOR_NAVY = RGBColor(10, 25, 47)          # #0A192F - Deep Primary
COLOR_BLUE = RGBColor(30, 58, 138)         # #1E3A8A - Tech Blue
COLOR_ACCENT = RGBColor(37, 99, 235)       # #2563EB - Electric Blue
COLOR_PURPLE = RGBColor(147, 51, 234)      # #9333EA - AgentX Purple
COLOR_CYAN = RGBColor(2, 132, 199)         # #0284C7 - Cyan
COLOR_GREEN = RGBColor(16, 185, 129)       # #10B981 - Success Emerald
COLOR_AMBER = RGBColor(245, 158, 11)       # #F59E0B - Warning Amber
COLOR_RED = RGBColor(239, 68, 68)          # #EF4444 - Problem Red
COLOR_BG_CARD = RGBColor(248, 250, 252)    # #F8FAFC - Card Background
COLOR_TEXT_MAIN = RGBColor(15, 23, 42)     # #0F172A - Main Text
COLOR_TEXT_MUTED = RGBColor(100, 116, 139) # #64748B - Muted Text
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_BORDER = RGBColor(226, 232, 240)


def create_header(slide, eyebrow_text, title_text):
    """Creates a clean, consistent header with the AgentX 2026 logo in top right (No marks)."""
    # Top decorative bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT
    top_bar.line.fill.background()

    # Header Box (Left Side)
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.5), Inches(1.1))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    # Eyebrow
    p_eye = tf.paragraphs[0]
    p_eye.text = eyebrow_text.upper()
    p_eye.font.size = Pt(11)
    p_eye.font.bold = True
    p_eye.font.color.rgb = COLOR_ACCENT

    # Title
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_NAVY
    p_title.space_before = Pt(2)

    # Logo in Top Right (Replaces Marks)
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Inches(10.6), Inches(0.35), width=Inches(2.0))


blank_slide_layout = prs.slide_layouts[6]

# ==============================================================================
# SLIDE 1: TITLE SLIDE
# ==============================================================================
slide1 = prs.slides.add_slide(blank_slide_layout)

# Background Accent Banner
bg_rect = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg_rect.fill.solid()
bg_rect.fill.fore_color.rgb = COLOR_NAVY
bg_rect.line.fill.background()

# Top Accent Line
top_line = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.18))
top_line.fill.solid()
top_line.fill.fore_color.rgb = COLOR_GREEN
top_line.line.fill.background()

# Main Title Card
title_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), Inches(11.333), Inches(5.5))
title_card.fill.solid()
title_card.fill.fore_color.rgb = RGBColor(15, 30, 58)
title_card.line.color.rgb = RGBColor(30, 58, 110)

# Logo inside Title Slide
if os.path.exists(LOGO_PATH):
    slide1.shapes.add_picture(LOGO_PATH, Inches(9.0), Inches(1.35), width=Inches(2.8))

t_box = slide1.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(7.5), Inches(4.8))
tf1 = t_box.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "ROUND-1 PRESENTATION"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_GREEN

p = tf1.add_paragraph()
p.text = "AI College Placement Assistant"
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Autonomous AI Agent for Student-Job Matching, Skill Gap Diagnosis, 7-Day Roadmaps & LaTeX Resume Tailoring"
p.font.size = Pt(15)
p.font.color.rgb = RGBColor(199, 210, 254)
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "───────────────────────────────────────────────────────────────────────────────────"
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(71, 85, 105)
p.space_before = Pt(12)

p = tf1.add_paragraph()
p.text = "Domain: AI & Campus Placement Automation   |   Cloud Database: Turso Cloud"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = COLOR_WHITE
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Team Members: Full-Stack Lead • AI Agent Specialist • Frontend & LaTeX Engineer • Cloud & QA Lead"
p.font.size = Pt(12)
p.font.color.rgb = RGBColor(148, 163, 184)
p.space_before = Pt(6)


# ==============================================================================
# SLIDE 2: THE PROBLEM
# ==============================================================================
slide2 = prs.slides.add_slide(blank_slide_layout)
create_header(slide2, "Problem Statement", "The Broken Campus Placement Lifecycle")

# Problem Summary Callout Box
prob_summary = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.9))
prob_summary.fill.solid()
prob_summary.fill.fore_color.rgb = RGBColor(254, 242, 242)
prob_summary.line.color.rgb = RGBColor(254, 202, 202)
ps_tf = prob_summary.text_frame
ps_tf.word_wrap = True
ps_p = ps_tf.paragraphs[0]
ps_p.text = '“Students miss top opportunities due to manual eligibility checks, unidentified skill gaps, and generic resumes, while placement cells struggle with fragmented data across 600+ candidates.”'
ps_p.font.size = Pt(13)
ps_p.font.bold = True
ps_p.font.color.rgb = RGBColor(153, 27, 27)
ps_p.alignment = PP_ALIGN.CENTER

# 3 Problem Pillar Cards
card_data = [
    ("📉 Blind Applications", "Missed Shortlists & Cutoffs", "Students apply without verifying multi-factor thresholds (CGPA, Branch, Year, Skills), causing high rejection rates.", RGBColor(254, 242, 242), RGBColor(239, 68, 68)),
    ("⏳ Last-Minute Panic", "No Structured Prep Roadmaps", "Lack of personalized day-by-day interview preparation schedules and company-specific mock technical questions.", RGBColor(255, 251, 235), RGBColor(245, 158, 11)),
    ("📄 Generic Resumes", "ATS Filtering & Zero Tailoring", "Submitting static, un-tailored resumes without STAR-method project blueprints matching job description keywords.", RGBColor(240, 253, 244), RGBColor(16, 185, 129))
]

for idx, (head, sub, desc, bg_col, accent_col) in enumerate(card_data):
    left = Inches(0.8 + idx * 4.0)
    card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(2.7), Inches(3.733), Inches(3.2))
    card.fill.solid()
    card.fill.fore_color.rgb = bg_col
    card.line.color.rgb = COLOR_BORDER

    tb = slide2.shapes.add_textbox(left + Inches(0.2), Inches(2.9), Inches(3.333), Inches(2.8))
    c_tf = tb.text_frame
    c_tf.word_wrap = True

    p = c_tf.paragraphs[0]
    p.text = head
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = accent_col

    p = c_tf.add_paragraph()
    p.text = sub
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY
    p.space_before = Pt(6)

    p = c_tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(10)

# Bottom Banner
bot_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.1), Inches(11.733), Inches(0.8))
bot_box.fill.solid()
bot_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
bot_box.line.color.rgb = RGBColor(199, 210, 254)
b_tf = bot_box.text_frame
b_p = b_tf.paragraphs[0]
b_p.text = "🔒 The Off-Campus Blindspot: Campus portals hide live market drives — students remain unaware of active hiring at Uber, Google, Stripe, etc."
b_p.font.size = Pt(12)
b_p.font.bold = True
b_p.font.color.rgb = COLOR_BLUE
b_p.alignment = PP_ALIGN.CENTER


# ==============================================================================
# SLIDE 3: OUR SOLUTION
# ==============================================================================
slide3 = prs.slides.add_slide(blank_slide_layout)
create_header(slide3, "Proposed Solution", "AI College Placement Assistant Platform")

features = [
    ("🎯 Automated Matching & Eligibility", "Computes multi-factor fit scores (0-100%) against academic cutoffs and skill taxonomies for 600+ students."),
    ("⚡ Autonomous Auto-Apply (LangGraph)", "Multi-step StateGraph evaluates fit, tailors screening answers, and submits applications with verification receipts."),
    ("🛠️ Zero-API-Key Skill Gap Planner", "Deterministic expert knowledge base generating 7-day day-by-day schedules & company mock questions 100% offline."),
    ("📄 ATS-Optimized LaTeX Resumes (.tex)", "Dynamically injects student data and STAR-format project blueprints into compilation-ready LaTeX templates.")
]

for idx, (title, desc) in enumerate(features):
    col = idx % 2
    row = idx // 2
    left = Inches(0.8 + col * 6.0)
    top = Inches(1.6 + row * 2.3)

    box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(5.733), Inches(2.0))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_BG_CARD
    box.line.color.rgb = COLOR_ACCENT

    tb = slide3.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), Inches(5.333), Inches(1.7))
    f_tf = tb.text_frame
    f_tf.word_wrap = True

    p = f_tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY

    p = f_tf.add_paragraph()
    p.text = desc
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(6)

# Bottom Workflow Strip
wf_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.65))
wf_box.fill.solid()
wf_box.fill.fore_color.rgb = COLOR_NAVY
wf_box.line.fill.background()
w_tf = wf_box.text_frame
w_p = w_tf.paragraphs[0]
w_p.text = "Workflow: Student Profile  ──►  Eligibility Engine  ──►  LangGraph Auto-Apply  ──►  7-Day Prep Roadmap  ──►  LaTeX Resume (.tex)"
w_p.font.size = Pt(11)
w_p.font.bold = True
w_p.font.color.rgb = COLOR_WHITE
w_p.alignment = PP_ALIGN.CENTER


# ==============================================================================
# SLIDE 4: TECHNICAL ARCHITECTURE (END-TO-END FLOW)
# ==============================================================================
slide4 = prs.slides.add_slide(blank_slide_layout)
create_header(slide4, "System Architecture", "End-to-End Architectural Data Flow")

# 4 Horizontal Flow Stage Cards with Connectors
flow_stages = [
    ("STAGE 1: INPUT & DATA", "📥 Raw Ingestion", "• 600+ Student Profiles\n• Turso Cloud Database\n• Live Internet Market Feeds\n• Company Job Schemas", COLOR_CYAN),
    ("STAGE 2: ORCHESTRATION", "⚙️ FastAPI Core", "• JWT Auth & RBAC\n• Pydantic v2 Validation\n• Jaccard Skill Parser\n• Academic Cutoff Filter", COLOR_BLUE),
    ("STAGE 3: AI AGENT CORE", "🤖 LangGraph Engine", "• Multi-Node StateGraph\n• Conditional Fit Routing\n• Offline Knowledge Graph\n• STAR Project Synthesizer", COLOR_PURPLE),
    ("STAGE 4: OUTPUTS", "🚀 Autonomous Action", "• Cryptographic Receipts\n• 7-Day Action Roadmaps\n• Company Mock Q&A\n• Compilation .tex Resume", COLOR_GREEN)
]

for idx, (stage_num, stage_title, items, border_color) in enumerate(flow_stages):
    left = Inches(0.8 + idx * 2.95)
    
    # Stage Container Card
    box = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, Inches(1.6), Inches(2.7), Inches(4.4))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_BG_CARD
    box.line.color.rgb = border_color
    box.line.width = Pt(2)

    # Top Header Strip
    top_strip = slide4.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, Inches(1.6), Inches(2.7), Inches(0.45))
    top_strip.fill.solid()
    top_strip.fill.fore_color.rgb = border_color
    top_strip.line.fill.background()
    s_tf = top_strip.text_frame
    sp = s_tf.paragraphs[0]
    sp.text = stage_num
    sp.font.size = Pt(10)
    sp.font.bold = True
    sp.font.color.rgb = COLOR_WHITE
    sp.alignment = PP_ALIGN.CENTER

    # Content Text Box
    tb = slide4.shapes.add_textbox(left + Inches(0.15), Inches(2.15), Inches(2.4), Inches(3.7))
    tf_f = tb.text_frame
    tf_f.word_wrap = True

    p = tf_f.paragraphs[0]
    p.text = stage_title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY

    p = tf_f.add_paragraph()
    p.text = items
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.line_spacing = 1.3
    p.space_before = Pt(8)

    # Add Arrow Connector between stages
    if idx < 3:
        arrow_left = left + Inches(2.72)
        arrow = slide4.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, arrow_left, Inches(3.6), Inches(0.2), Inches(0.3))
        arrow.fill.solid()
        arrow.fill.fore_color.rgb = COLOR_ACCENT
        arrow.line.fill.background()

# Bottom Flow Callout Box
arch_note = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.2), Inches(11.733), Inches(0.75))
arch_note.fill.solid()
arch_note.fill.fore_color.rgb = RGBColor(240, 253, 244)
arch_note.line.color.rgb = RGBColor(187, 247, 208)
an_tf = arch_note.text_frame
an_p = an_tf.paragraphs[0]
an_p.text = "⚡ Architectural Flow: Data Ingestion ──► FastAPI RBAC Filtering ──► LangGraph State Machine ──► Auto-Apply & LaTeX Output"
an_p.font.size = Pt(11)
an_p.font.bold = True
an_p.font.color.rgb = RGBColor(22, 101, 52)
an_p.alignment = PP_ALIGN.CENTER


# ==============================================================================
# SLIDE 5: TECH STACK — "HOW WE BUILT IT"
# ==============================================================================
slide5 = prs.slides.add_slide(blank_slide_layout)
create_header(slide5, "Implementation", "How We Built It — Technology Stack")

rows = [
    ("Frontend Layer", "React 19.x, Vite 6, React Router DOM, Custom Glassmorphic CSS3 System"),
    ("Backend API Server", "Python 3.13, FastAPI (Async), Uvicorn ASGI Server, Pydantic v2"),
    ("Agentic AI & NLP", "LangGraph (StateGraph), Hugging Face Hub, Knowledge Graph Expert System"),
    ("Database & Cloud", "Turso Cloud DB (Serverless libSQL over HTTPS), SQLite 3, SQLAlchemy 2.0"),
    ("Security & RBAC", "PyJWT (JSON Web Tokens), Passlib (Salted Bcrypt), Granular FastAPI Roles"),
    ("Document Engineering", "LaTeX (.tex dynamic markup generation, character sanitization engine)"),
    ("Testing & QA", "Pytest 9.x (15 Unit & Integration Test Suites Passing in 2.7s)")
]

top_offset = 1.6
for idx, (cat, stack) in enumerate(rows):
    bg_c = COLOR_BG_CARD if idx % 2 == 0 else RGBColor(255, 255, 255)
    row_box = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(top_offset + idx * 0.65), Inches(11.733), Inches(0.6))
    row_box.fill.solid()
    row_box.fill.fore_color.rgb = bg_c
    row_box.line.color.rgb = COLOR_BORDER

    # Left Label
    lbl_box = slide5.shapes.add_textbox(Inches(0.9), Inches(top_offset + idx * 0.65 + 0.05), Inches(2.8), Inches(0.5))
    ltf = lbl_box.text_frame
    lp = ltf.paragraphs[0]
    lp.text = cat
    lp.font.size = Pt(12)
    lp.font.bold = True
    lp.font.color.rgb = COLOR_NAVY

    # Right Content
    val_box = slide5.shapes.add_textbox(Inches(3.8), Inches(top_offset + idx * 0.65 + 0.05), Inches(8.5), Inches(0.5))
    vtf = val_box.text_frame
    vp = vtf.paragraphs[0]
    vp.text = stack
    vp.font.size = Pt(12)
    vp.font.color.rgb = COLOR_TEXT_MAIN

# Bottom Testing Badge
test_box = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.65))
test_box.fill.solid()
test_box.fill.fore_color.rgb = RGBColor(238, 242, 255)
test_box.line.color.rgb = RGBColor(199, 210, 254)
t_tf = test_box.text_frame
tp = t_tf.paragraphs[0]
tp.text = "✓ Quality Assured: 15/15 automated tests passing with pytest (LangGraph workflows, eligibility, & LaTeX engine)."
tp.font.size = Pt(11)
tp.font.bold = True
tp.font.color.rgb = COLOR_BLUE
tp.alignment = PP_ALIGN.CENTER


# ==============================================================================
# SLIDE 6: FUTURE SCOPE & TEAM WORK
# ==============================================================================
slide6 = prs.slides.add_slide(blank_slide_layout)
create_header(slide6, "Scalability & Collaboration", "Future Enhancements & Team Contribution")

# Left Column: Future Scope
left_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.5))
left_card.fill.solid()
left_card.fill.fore_color.rgb = COLOR_BG_CARD
left_card.line.color.rgb = COLOR_ACCENT

ltb = slide6.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.1))
ltf = ltb.text_frame
ltf.word_wrap = True

p = ltf.paragraphs[0]
p.text = "🚀 Future Enhancements Roadmap"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

roadmap = [
    ("🎥 AI Video Mock Interview Copilot", "Real-time speech rate, filler-word detection, and facial confidence analysis."),
    ("📱 WhatsApp & Telegram Push Alerts", "Instant notification dispatch when high-fit on-campus & off-campus drives open."),
    ("🌐 Multi-College Placement Grid", "Inter-collegiate talent pooling for shared pooled campus placement drives.")
]

for r_t, r_d in roadmap:
    p = ltf.add_paragraph()
    p.text = r_t
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLUE
    p.space_before = Pt(10)

    p = ltf.add_paragraph()
    p.text = r_d
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(2)

# Right Column: Team Roles
right_card = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.6), Inches(5.733), Inches(4.5))
right_card.fill.solid()
right_card.fill.fore_color.rgb = COLOR_GREEN

rtb = slide6.shapes.add_textbox(Inches(7.0), Inches(1.8), Inches(5.333), Inches(4.1))
rtf = rtb.text_frame
rtf.word_wrap = True

p = rtf.paragraphs[0]
p.text = "👥 Team Roles & Ownership"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = COLOR_NAVY

team_roles = [
    ("Full-Stack & Backend Lead", "FastAPI architecture, JWT RBAC security, REST endpoints"),
    ("AI Agent & NLP Specialist", "LangGraph StateGraph, conditional edges, matching engine"),
    ("Frontend UI/UX & LaTeX Lead", "React 19 dashboards, LeetCode arena, .tex generator"),
    ("Cloud DB & Testing Lead", "Turso Cloud DB sync, CSV ETL ingestion, Pytest test suites")
]

for role, desc in team_roles:
    p = rtf.add_paragraph()
    p.text = f"• {role}"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY
    p.space_before = Pt(8)

    p = rtf.add_paragraph()
    p.text = f"  {desc}"
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(1)

# Bottom Collab Note
collab_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.3), Inches(11.733), Inches(0.65))
collab_box.fill.solid()
collab_box.fill.fore_color.rgb = COLOR_NAVY
collab_box.line.fill.background()
c_tf = collab_box.text_frame
cp = c_tf.paragraphs[0]
cp.text = "Collaboration: Agile Git workflows, modular branch strategy, and end-to-end integration testing."
cp.font.size = Pt(11)
cp.font.bold = True
cp.font.color.rgb = COLOR_WHITE
cp.alignment = PP_ALIGN.CENTER


# ==============================================================================
# SLIDE 7: LIVE DEMO
# ==============================================================================
slide7 = prs.slides.add_slide(blank_slide_layout)
create_header(slide7, "Practical Demonstration", "Live Demo & Real-World Use Case")

# Persona Box
persona = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), Inches(11.733), Inches(0.8))
persona.fill.solid()
persona.fill.fore_color.rgb = RGBColor(238, 242, 255)
persona.line.color.rgb = COLOR_ACCENT
p_tf = persona.text_frame
p_p = p_tf.paragraphs[0]
p_p.text = "🧑‍💻 Real-World Persona: Mohammad Rafi — Final-Year CSE Student (CGPA: 9.60, Target: Top Tech SDE Roles)"
p_p.font.size = Pt(12)
p_p.font.bold = True
p_p.font.color.rgb = COLOR_BLUE
p_p.alignment = PP_ALIGN.CENTER

# 6 Demo Steps Grid
steps = [
    ("1. Student Login", "Instant profile auto-hydration (Branch, CGPA, verified skill taxonomy)."),
    ("2. Opportunity Matching", "Composite fit percentage calculated in real time (Google SDE: 85% Match)."),
    ("3. 7-Day AI Prep Plan", "Exact skill gap breakdown, day-by-day roadmap, and company mock questions."),
    ("4. Agent Copilot Chat", "Instant 3-day crash course re-planning with zero external API key costs."),
    ("5. LangGraph Auto-Apply", "⚡ Auto-apply with cryptographic confirmation receipt (APP-XXXXXXXX)."),
    ("6. Tailored LaTeX Resume", "📄 1-click generation of compilation-ready .tex resume with STAR projects.")
]

for idx, (st_t, st_d) in enumerate(steps):
    col = idx % 3
    row = idx // 3
    left = Inches(0.8 + col * 4.0)
    top = Inches(2.6 + row * 1.6)

    box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.733), Inches(1.4))
    box.fill.solid()
    box.fill.fore_color.rgb = COLOR_BG_CARD
    box.line.color.rgb = COLOR_BORDER

    tb = slide7.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), Inches(3.433), Inches(1.2))
    stf = tb.text_frame
    stf.word_wrap = True

    p = stf.paragraphs[0]
    p.text = st_t
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_NAVY

    p = stf.add_paragraph()
    p.text = st_d
    p.font.size = Pt(10.5)
    p.font.color.rgb = COLOR_TEXT_MAIN
    p.space_before = Pt(3)

# Bottom WOW Factor Strip
wow_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.95))
wow_box.fill.solid()
wow_box.fill.fore_color.rgb = RGBColor(240, 253, 244)
wow_box.line.color.rgb = COLOR_GREEN
w_tf = wow_box.text_frame
w_p = w_tf.paragraphs[0]
w_p.text = "🌟 WOW Factors to Highlight during Demo:"
w_p.font.size = Pt(11)
w_p.font.bold = True
w_p.font.color.rgb = RGBColor(22, 101, 52)

w_p2 = w_tf.add_paragraph()
w_p2.text = "1. 100% Zero-API-Key Offline Reliability ($0 cost)   |   2. Autonomous LangGraph Auto-Apply   |   3. Dynamic LaTeX Resume Tailoring"
w_p2.font.size = Pt(11)
w_p2.font.color.rgb = RGBColor(22, 101, 52)
w_p2.space_before = Pt(2)


# Save presentation
output_path = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "AI_College_Placement_Assistant_Round1_Presentation.pptx"))
prs.save(output_path)
print(f"[+] Updated PowerPoint Presentation with AgentX Logo and Flow Architecture saved to:\n    {output_path}")
